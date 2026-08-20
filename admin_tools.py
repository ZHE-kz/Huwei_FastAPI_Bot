import csv
import hmac
import io
import re

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from openpyxl import load_workbook
from pydantic import BaseModel

from config import ADMIN_TOKEN
from firestore_client import get_db
from puzzle_core import PUZZLE_COLLECTION, force_clear_puzzle_cache, load_puzzle_config, normalize_stage

router = APIRouter(prefix="/admin/puzzles")
security = HTTPBasic(auto_error=False)
IMPORT_COLUMNS = [
    "stage_id",
    "agent_intro",
    "vision_target",
    "puzzle_hint",
    "puzzle_answer",
    "success_text",
    "next_stage_id",
    "requires_image",
]
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 200


def require_admin(credentials: HTTPBasicCredentials = Depends(security)):
    valid = (
        credentials
        and hmac.compare_digest(credentials.username, "admin")
        and hmac.compare_digest(credentials.password, ADMIN_TOKEN)
    )
    if not ADMIN_TOKEN:
        raise HTTPException(503, "ADMIN_TOKEN 尚未設定")
    if not valid:
        raise HTTPException(401, "Unauthorized", headers={"WWW-Authenticate": "Basic"})


class PuzzleSettings(BaseModel):
    start_command: str
    start_stage: str


class PuzzleStage(BaseModel):
    agent_intro: str
    vision_target: str = ""
    puzzle_hint: str
    puzzle_answer: str
    success_text: str
    next_stage_id: str = ""
    requires_image: bool = False


def parse_bool(value):
    return str(value or "").strip().casefold() in {"1", "true", "yes", "y", "是", "需要"}


def read_spreadsheet(filename, content):
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == ".csv":
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        rows = list(reader)
    elif suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
    else:
        raise HTTPException(400, "僅支援 .csv 或 .xlsx")

    if not rows:
        raise HTTPException(400, "試算表是空的")

    headers = [str(value or "").strip() for value in rows[0]]
    nonempty_headers = [header for header in headers if header]
    if len(nonempty_headers) != len(set(nonempty_headers)):
        raise HTTPException(400, "欄位名稱不可重複")
    if "stage_id" not in headers or "puzzle_answer" not in headers:
        raise HTTPException(400, "缺少必要欄位：stage_id、puzzle_answer")

    return [
        {header: value for header, value in zip(headers, row) if header}
        for row in rows[1:]
        if any(value not in (None, "") for value in row)
    ]


def validate_import_rows(rows):
    if not rows:
        raise HTTPException(400, "沒有可匯入的題目")
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(400, f"一次最多匯入 {MAX_IMPORT_ROWS} 題")

    stages = []
    errors = []
    for row_number, row in enumerate(rows, start=2):
        stage_id = str(row.get("stage_id") or "").strip()
        answer = str(row.get("puzzle_answer") or "").strip()
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,50}", stage_id):
            errors.append(f"第 {row_number} 列：stage_id 格式錯誤")
            continue
        if not answer:
            errors.append(f"第 {row_number} 列：puzzle_answer 不可空白")
            continue

        stages.append(
            normalize_stage(
                stage_id,
                {
                    "agent_intro": str(row.get("agent_intro") or "").strip(),
                    "vision_target": str(row.get("vision_target") or "").strip(),
                    "puzzle_hint": str(row.get("puzzle_hint") or "").strip(),
                    "puzzle_answer": answer,
                    "success_text": str(row.get("success_text") or "").strip(),
                    "next_stage_id": str(row.get("next_stage_id") or "").strip(),
                    "requires_image": parse_bool(row.get("requires_image")),
                },
            )
        )

    if errors:
        raise HTTPException(400, {"errors": errors[:20]})
    return stages


@router.get("", response_class=HTMLResponse, dependencies=[Depends(require_admin)])
async def puzzle_admin_page():
    return HTMLResponse(ADMIN_HTML)


@router.get("/template.csv", dependencies=[Depends(require_admin)])
async def puzzle_template():
    sample = [
        "STAGE_01",
        "第一關開始，請依提示完成任務。",
        "虎尾糖廠煙囪",
        "煙囪主要是什麼顏色？",
        "紅色",
        "答對了！",
        "STAGE_02",
        "true",
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(IMPORT_COLUMNS)
    writer.writerow(sample)
    return PlainTextResponse(
        "\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="puzzle-template.csv"'},
    )


@router.get("/api", dependencies=[Depends(require_admin)])
async def list_puzzles():
    return await load_puzzle_config()


@router.post("/api/import", dependencies=[Depends(require_admin)])
async def import_puzzles(file: UploadFile = File(...)):
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "檔案不可超過 5 MB")

    try:
        stages = validate_import_rows(read_spreadsheet(file.filename or "", content))
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(400, f"無法讀取試算表：{error}") from error
    database = get_db()
    batch = database.batch()
    for stage in stages:
        batch.set(database.collection(PUZZLE_COLLECTION).document(stage["stage_id"]), stage)
    batch.commit()
    force_clear_puzzle_cache()
    return {"ok": True, "imported": len(stages)}


@router.put("/api/settings", dependencies=[Depends(require_admin)])
async def update_settings(settings: PuzzleSettings):
    if not settings.start_command.strip() or not settings.start_stage.strip():
        raise HTTPException(400, "啟動指令與起始關卡不可空白")
    get_db().collection(PUZZLE_COLLECTION).document("_settings").set(settings.model_dump())
    force_clear_puzzle_cache()
    return {"ok": True}


@router.put("/api/stages/{stage_id}", dependencies=[Depends(require_admin)])
async def update_stage(stage_id: str, stage: PuzzleStage):
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,50}", stage_id):
        raise HTTPException(400, "關卡 ID 只能使用英數、底線與連字號")
    data = normalize_stage(stage_id, stage.model_dump())
    get_db().collection(PUZZLE_COLLECTION).document(stage_id).set(data)
    force_clear_puzzle_cache()
    return {"ok": True, "stage": data}


ADMIN_HTML = """<!doctype html>
<html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>解謎題庫管理</title><style>
body{font:16px system-ui;max-width:980px;margin:auto;padding:24px;background:#f5f7f5;color:#18311f}h1{margin-top:0}
section{background:white;padding:20px;margin:16px 0;border-radius:12px;box-shadow:0 2px 12px #0001}label{display:block;margin:10px 0 4px}
input,textarea{box-sizing:border-box;width:100%;padding:10px;border:1px solid #b8c7bb;border-radius:7px}textarea{min-height:76px}button,.button{display:inline-block;margin-top:14px;padding:10px 16px;border:0;border-radius:7px;background:#1d7a3b;color:white;cursor:pointer;text-decoration:none}.grid{display:grid;grid-template-columns:1fr 1fr;gap:12px}.stage{border-top:1px solid #ddd;padding:12px 0}.muted{color:#607067}.status{margin-left:10px}@media(max-width:650px){.grid{grid-template-columns:1fr}}
</style></head><body><h1>解謎題庫管理</h1><p class="muted">使用者輸入啟動指令後，從起始關卡開始。修改會直接寫入 Firestore。</p>
<section><h2>遊戲設定</h2><form id="settings"><div class="grid"><div><label>啟動指令</label><input name="start_command" required></div><div><label>起始關卡</label><input name="start_stage" required></div></div><button>儲存設定</button></form></section>
<section><h2>試算表匯入</h2><p class="muted">支援 CSV、XLSX，最多 5 MB／200 題。同 ID 關卡會覆寫，其他既有關卡不刪除。</p><form id="upload"><input type="file" name="file" accept=".csv,.xlsx" required><button>上傳並匯入</button><span id="upload-status" class="status"></span></form><a class="button" href="/admin/puzzles/template.csv">下載 CSV 範本</a></section>
<section><h2>新增／修改關卡</h2><form id="stage"><div class="grid"><div><label>關卡 ID</label><input name="stage_id" placeholder="STAGE_01" required></div><div><label>下一關 ID</label><input name="next_stage_id"></div></div><label>開場文字</label><textarea name="agent_intro" required></textarea><label>圖片辨識目標</label><textarea name="vision_target"></textarea><label><input style="width:auto" type="checkbox" name="requires_image"> 此關需要先上傳照片</label><label>題目／提示</label><textarea name="puzzle_hint" required></textarea><label>正確答案</label><input name="puzzle_answer" required><label>答對訊息</label><textarea name="success_text" required></textarea><button>儲存關卡</button></form></section>
<section><h2>目前關卡</h2><div id="list"></div></section><script>
const settings=document.querySelector('#settings'),stage=document.querySelector('#stage'),upload=document.querySelector('#upload'),list=document.querySelector('#list'),uploadStatus=document.querySelector('#upload-status');
async function api(path='',options={}){const headers=options.body instanceof FormData?{}:{'Content-Type':'application/json'};const r=await fetch('/admin/puzzles/api'+path,{...options,headers:{...headers,...options.headers}});if(!r.ok)throw new Error(await r.text());return r.json()}
async function load(){const data=await api();settings.start_command.value=data.settings.start_command;settings.start_stage.value=data.settings.start_stage;list.replaceChildren(...Object.entries(data.stages).sort().map(([id,s])=>{const row=document.createElement('div');row.className='stage';const title=document.createElement('strong');title.textContent=id+' — '+s.puzzle_hint;const edit=document.createElement('button');edit.textContent='編輯';edit.onclick=()=>fill(id,s);row.append(title,document.createElement('br'),edit);return row}))}
function fill(id,s){stage.stage_id.value=id;for(const [key,value] of Object.entries(s)){if(stage[key])stage[key].type==='checkbox'?stage[key].checked=value:stage[key].value=value??''}stage.scrollIntoView({behavior:'smooth'})}
settings.onsubmit=async e=>{e.preventDefault();await api('/settings',{method:'PUT',body:JSON.stringify(Object.fromEntries(new FormData(settings)))});alert('設定已儲存')};
upload.onsubmit=async e=>{e.preventDefault();uploadStatus.textContent='匯入中…';try{const result=await api('/import',{method:'POST',body:new FormData(upload)});uploadStatus.textContent='已匯入 '+result.imported+' 題';upload.reset();await load()}catch(error){uploadStatus.textContent='匯入失敗：'+error.message}};
stage.onsubmit=async e=>{e.preventDefault();const data=Object.fromEntries(new FormData(stage));const id=data.stage_id;delete data.stage_id;data.requires_image=stage.requires_image.checked;await api('/stages/'+encodeURIComponent(id),{method:'PUT',body:JSON.stringify(data)});alert('關卡已儲存');await load()};load();
</script></body></html>"""